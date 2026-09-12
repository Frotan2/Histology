#!/usr/bin/env python3
"""Build the Question Bank DOCX and EPUB from MD source."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
SRC = ROOT / 'Essential_Histology_Question_Bank.md'
OUT_DOCX = ROOT / 'Essential_Histology_Question_Bank.docx'
OUT_EPUB = ROOT / 'Essential_Histology_Question_Bank.epub'


def build_docx(md_path: str, docx_path: str):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    # Style
    style = doc.styles['Normal']
    style.font.name = 'DejaVu Serif'
    style.font.size = Pt(11)
    # Title
    title = doc.add_heading('Essential Histology: Question Bank', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph('A Concept-Based Guide for PGME — Assessment Companion')
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(12)
    sub.runs[0].italic = True
    doc.add_paragraph('Edition: Essential Histology 1405 — Definitive Edition (v3)')
    doc.add_paragraph('Primary reference: Junqueira 17th Ed. (Mescher, McGraw Hill)')
    doc.add_paragraph('Official syllabus: 16 chapters mapped to Junqueira 1, 2, 3, 4, 5, 6, 9, 11, 12, 13, 14, 15, 16, 17, 18, 20.')
    doc.add_paragraph('Each question maps to chapter → content unit → learning objective → answer → explanation.')
    doc.add_paragraph()
    # Parse MD: ## for chapter, ### for part, **Q.** for question
    text = Path(md_path).read_text()
    lines = text.split('\n')
    in_q = False
    for line in lines:
        if line.startswith('# Essential Histology: Question Bank'):
            continue
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('#### ') or line.startswith('##### '):
            doc.add_heading(line.lstrip('# '), level=3)
        elif line.startswith('**Q'):
            doc.add_paragraph().add_run(line).bold = True
        elif line.startswith('**Answer'):
            doc.add_paragraph().add_run(line).bold = True
        elif line.startswith('**Explanation'):
            doc.add_paragraph().add_run(line).bold = True
        elif line.startswith('**Decisive feature'):
            doc.add_paragraph().add_run(line).bold = True
        elif line.startswith('**Learning Objective'):
            doc.add_paragraph().add_run(line).bold = True
        elif line.startswith('---'):
            doc.add_paragraph()
        elif line.strip():
            doc.add_paragraph(line)
    doc.save(docx_path)
    print(f'  DOCX: {Path(docx_path).stat().st_size:,} bytes')


def build_epub(md_path: str, epub_path: str):
    from ebooklib import epub
    md = Path(md_path).read_text()
    book = epub.EpubBook()
    book.set_identifier('essential-histology-qb-v3')
    book.set_title('Essential Histology: Question Bank — Definitive Edition')
    book.set_language('en')
    book.add_author('AREMS-HY Academic Series')
    # Title page
    tp = epub.EpubHtml(title='Title Page', file_name='title.xhtml', lang='en')
    tp.content = '<h1>Essential Histology: Question Bank</h1><h2>A Concept-Based Guide for PGME — Assessment Companion</h2><p><i>Edition: Essential Histology 1405 — Definitive Edition (v3)</i></p><p>Primary reference: Junqueira 17th Ed.</p>'
    book.add_item(tp)
    # Split by ## (chapters)
    chapters_md = []
    current = []
    for line in md.split('\n'):
        if line.startswith('## ') and current:
            chapters_md.append(('\n'.join(current)))
            current = [line]
        else:
            current.append(line)
    if current:
        chapters_md.append('\n'.join(current))
    chapters_md = [c for c in chapters_md if c.startswith('## ')]
    items = [tp]
    for i, ch_md in enumerate(chapters_md):
        title_line = ch_md.split('\n', 1)[0]
        title = title_line.lstrip('# ').strip()
        # HTML conversion simple
        body = ch_md.split('\n', 1)[1] if '\n' in ch_md else ''
        html = '<h2>' + title + '</h2>'
        for ln in body.split('\n'):
            if ln.startswith('### '):
                html += '<h3>' + ln[4:].strip() + '</h3>'
            elif ln.startswith('**Q'):
                html += '<p><b>' + ln.strip() + '</b></p>'
            elif ln.startswith('**Answer'):
                html += '<p><b>' + ln.strip() + '</b></p>'
            elif ln.startswith('**Explanation'):
                html += '<p><b>' + ln.strip() + '</b></p>'
            elif ln.startswith('**Decisive'):
                html += '<p><b>' + ln.strip() + '</b></p>'
            elif ln.startswith('**Learning'):
                html += '<p><b>' + ln.strip() + '</b></p>'
            elif ln.strip() == '---':
                html += '<hr/>'
            elif ln.strip():
                html += '<p>' + ln.strip() + '</p>'
        item = epub.EpubHtml(title=title, file_name=f'ch_{i+1:02d}.xhtml', lang='en')
        item.content = html
        book.add_item(item)
        items.append(item)
    book.toc = tuple(items)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ['nav'] + items
    epub.write_epub(epub_path, book)
    print(f'  EPUB: {Path(epub_path).stat().st_size:,} bytes')


if __name__ == '__main__':
    build_docx(str(SRC), str(OUT_DOCX))
    build_epub(str(SRC), str(OUT_EPUB))
