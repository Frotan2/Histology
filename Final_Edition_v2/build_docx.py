#!/usr/bin/env python3
"""
Build the final DOCX from canonical markdown source.
"""
import os
import re
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Source and output paths
SRC = os.path.join(os.path.dirname(__file__), 'source', 'Essential_Histology_Definitive_Final.md')
IMG_DIR = os.path.join(os.path.dirname(__file__), 'source', 'images')
OUT = os.path.join(os.path.dirname(__file__), 'Essential_Histology_v2.docx')

# Color palette
NAVY = RGBColor(0x0A, 0x3D, 0x62)
DARK_GRAY = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
LIGHT_NAVY = RGBColor(0x2C, 0x3E, 0x50)

# Font setup
BODY_FONT = 'Georgia'  # Serif, widely available, similar to Garamond
HEAD_FONT = 'Calibri'  # Sans, widely available

def add_page_break(paragraph):
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)

def set_paragraph_spacing(paragraph, before=0, after=6, line_spacing=1.15):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing

def add_run_with_format(paragraph, text, bold=False, italic=False, size=11, color=DARK_GRAY, font=BODY_FONT):
    if not text:
        return
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    return run

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '999999')
        tcBorders.append(border)
    tcPr.append(tcBorders)

def parse_inline(text, paragraph, base_size=11, base_color=DARK_GRAY, base_font=BODY_FONT):
    """Parse inline markdown (bold, italic, code) into runs."""
    # Handle bold-italic patterns
    pattern = re.compile(r'(\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            chunk = text[pos:m.start()]
            if chunk:
                add_run_with_format(paragraph, chunk, size=base_size, color=base_color, font=base_font)
        if m.group(3):  # **bold**
            add_run_with_format(paragraph, m.group(3), bold=True, size=base_size, color=base_color, font=base_font)
        elif m.group(4):  # *italic*
            add_run_with_format(paragraph, m.group(4), italic=True, size=base_size, color=base_color, font=base_font)
        elif m.group(2):  # ***bold-italic***
            add_run_with_format(paragraph, m.group(2), bold=True, italic=True, size=base_size, color=base_color, font=base_font)
        elif m.group(5):  # `code`
            add_run_with_format(paragraph, m.group(5), size=base_size - 1, color=base_color, font='Consolas')
        pos = m.end()
    if pos < len(text):
        add_run_with_format(paragraph, text[pos:], size=base_size, color=base_color, font=base_font)

def add_image(paragraph, image_path, width_inches=6.0):
    """Add an image, scaling to fit the page width."""
    if not os.path.exists(image_path):
        return False
    try:
        run = paragraph.add_run()
        run.add_picture(image_path, width=Inches(width_inches))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return True
    except Exception as e:
        print(f'Error adding image {image_path}: {e}')
        return False

def render_table(doc, lines, start_idx):
    """Render a markdown table starting at lines[start_idx]."""
    # Parse table lines
    rows = []
    i = start_idx
    while i < len(lines) and lines[i].strip().startswith('|'):
        if re.match(r'^\|[\s\-\|:]+\|?$', lines[i].strip()):
            i += 1
            continue
        cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
        rows.append(cells)
        i += 1

    if not rows:
        return i

    header = rows[0]
    data_rows = rows[1:]

    # Normalize column count
    max_cols = max(len(r) for r in rows)
    header = header + [''] * (max_cols - len(header))
    data_rows = [r + [''] * (max_cols - len(r)) for r in data_rows]

    table = doc.add_table(rows=1 + len(data_rows), cols=max_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Header row
    for j, cell_text in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ''
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, before=2, after=2)
        parse_inline(cell_text, p, base_size=10, base_font=BODY_FONT)
        for run in p.runs:
            run.bold = True
        set_cell_shading(cell, 'EAEDED')
        set_cell_borders(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Data rows
    for ri, row in enumerate(data_rows):
        for j, cell_text in enumerate(row):
            cell = table.rows[ri + 1].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, before=1, after=1)
            parse_inline(cell_text, p, base_size=9, base_font=BODY_FONT)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    # Add space after table
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=4, after=4)
    return i

def process_markdown(doc, lines, skip_front_matter=False):
    """Process markdown lines into docx elements."""
    i = 0
    while i < len(lines):
        line = lines[i]

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # Page break
        if line.strip() == '\\pagebreak':
            p = doc.add_paragraph()
            add_page_break(p)
            i += 1
            continue

        # Horizontal rule
        if line.strip() == '---':
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=4, after=4)
            i += 1
            continue

        # Heading 1
        if line.startswith('# '):
            text = line[2:].strip()
            # Page break before chapter
            if 'Chapter' in text or 'Final' in text or 'Examination' in text or 'Recognition' in text:
                p = doc.add_paragraph()
                add_page_break(p)
            p = doc.add_heading('', level=1)
            set_paragraph_spacing(p, before=12, after=12, line_spacing=1.2)
            parse_inline(text, p, base_size=22, base_color=NAVY, base_font=HEAD_FONT)
            for run in p.runs:
                run.bold = True
            i += 1
            continue

        # Heading 2
        if line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_heading('', level=2)
            set_paragraph_spacing(p, before=12, after=6, line_spacing=1.2)
            parse_inline(text, p, base_size=15, base_color=ACCENT, base_font=HEAD_FONT)
            for run in p.runs:
                run.bold = True
            i += 1
            continue

        # Heading 3
        if line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_heading('', level=3)
            set_paragraph_spacing(p, before=8, after=4, line_spacing=1.15)
            parse_inline(text, p, base_size=12, base_color=LIGHT_NAVY, base_font=HEAD_FONT)
            for run in p.runs:
                run.bold = True
            i += 1
            continue

        # Heading 4
        if line.startswith('#### '):
            text = line[5:].strip()
            p = doc.add_heading('', level=4)
            set_paragraph_spacing(p, before=6, after=3, line_spacing=1.15)
            parse_inline(text, p, base_size=11, base_color=DARK_GRAY, base_font=HEAD_FONT)
            for run in p.runs:
                run.bold = True
                run.italic = True
            i += 1
            continue

        # Image
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if img_match:
            alt = img_match.group(1)
            src = img_match.group(2)
            image_path = os.path.join(os.path.dirname(SRC), src) if not os.path.isabs(src) else src
            if not os.path.exists(image_path):
                # Try relative to IMG_DIR
                fname = os.path.basename(src)
                image_path = os.path.join(IMG_DIR, fname)
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=8, after=4)
            if not add_image(p, image_path, width_inches=6.0):
                # Fallback: just show alt text
                add_run_with_format(p, f'[Image: {alt}]', italic=True, size=10, color=LIGHT_NAVY)
            i += 1
            continue

        # Table
        if line.strip().startswith('|'):
            i = render_table(doc, lines, i)
            continue

        # Bullet list
        if re.match(r'^\s*-\s', line):
            text = re.sub(r'^\s*-\s+', '', line)
            p = doc.add_paragraph(style='List Bullet')
            set_paragraph_spacing(p, before=2, after=2)
            parse_inline(text, p, base_size=11)
            i += 1
            continue

        # Numbered list (e.g. "1. text")
        num_match = re.match(r'^(\d+)\.\s+(.*)', line)
        if num_match and not line.startswith('#'):
            num = num_match.group(1)
            text = num_match.group(2)
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=2, after=2)
            add_run_with_format(p, f'{num}. ', bold=True, size=11)
            parse_inline(text, p, base_size=11)
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=2, after=4)
        parse_inline(line, p, base_size=11)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i += 1


def build_docx():
    """Main function to build DOCX from canonical markdown."""
    doc = Document()

    # Configure styles
    style = doc.styles['Normal']
    style.font.name = BODY_FONT
    style.font.size = Pt(11)

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Set headers
    for section in doc.sections:
        header = section.header
        h_para = header.paragraphs[0]
        h_para.text = ''
        h_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run_with_format(h_para, 'Essential Histology  —  Definitive Edition', italic=True, size=9, color=LIGHT_NAVY)

        footer = section.footer
        f_para = footer.paragraphs[0]
        f_para.text = ''
        f_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Page number field
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = 'PAGE'
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        run = f_para.add_run()
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        for r in f_para.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = LIGHT_NAVY

    # Add title page (special handling)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=120, after=12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run_with_format(p, 'Essential Histology', bold=True, size=36, color=NAVY, font=HEAD_FONT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=24)
    add_run_with_format(p, 'A Concept-Based Guide for PGME', italic=True, size=18, color=ACCENT, font=HEAD_FONT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=12, after=12)
    add_run_with_format(p, 'Understanding the structure. Seeing the pattern.', italic=True, size=12, color=DARK_GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=12)
    add_run_with_format(p, 'Distinguishing the alternatives. Retrieving the logic.', italic=True, size=12, color=DARK_GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=24, after=4)
    add_run_with_format(p, 'Definitive Edition', bold=True, size=14, color=NAVY, font=HEAD_FONT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=24)
    add_run_with_format(p, 'AREMS-HY Academic Series', italic=True, size=12, color=DARK_GRAY)

    # Reference and scope
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=48, after=4)
    add_run_with_format(p, 'Primary scientific foundation', italic=True, size=10, color=LIGHT_NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=4)
    add_run_with_format(p, 'Mescher AL. Junqueira\'s Basic Histology: Text and Atlas. 17th ed. McGraw Hill.', size=10, color=DARK_GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=12, after=4)
    add_run_with_format(p, 'Official 16-Chapter Syllabus', italic=True, size=10, color=LIGHT_NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=4)
    add_run_with_format(p, 'Aligned with Junqueira 17th Chapters 1, 2, 3, 4, 5, 6, 9, 11, 12, 13, 14, 15, 16, 17, 18, 20', size=10, color=DARK_GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=12, after=4)
    add_run_with_format(p, 'Built for the Afghanistan 1405 Specialty Examination', size=10, color=DARK_GRAY)

    # Page break before front matter
    p = doc.add_paragraph()
    add_page_break(p)

    # Read and process markdown (skip front matter title page from MD, we already added it)
    with open(SRC) as f:
        md_text = f.read()

    lines = md_text.split('\n')
    # Find the line with "# Essential Histology" (the front matter title)
    # Skip until we hit the "## Contents" section
    start = 0
    for i, line in enumerate(lines):
        if line.startswith('# Essential Histology'):
            start = i + 1
            break
    # Skip until we hit "## Contents"
    found_contents = False
    for i in range(start, len(lines)):
        if lines[i].startswith('## Contents'):
            start = i
            found_contents = True
            break
    if not found_contents:
        start = 0

    # Process from start
    process_markdown(doc, lines[start:], skip_front_matter=True)

    # Save
    doc.save(OUT)
    print(f'Wrote {OUT}: {os.path.getsize(OUT)} bytes')


if __name__ == '__main__':
    build_docx()
